"""文档导入 API 测试。"""
import io
import os
from pathlib import Path

from fastapi.testclient import TestClient


def _create_kb(client: TestClient, name: str = "测试知识库") -> dict:
    res = client.post("/api/v1/knowledge-bases", json={"name": name})
    assert res.status_code == 200
    return res.json()["data"]


# ---------- Markdown 导入 ----------


def test_upload_markdown(client: TestClient) -> None:
    """上传 .md 文件应成功创建 Wiki 页面。"""
    kb = _create_kb(client)
    content = "# 测试标题\n\n正文内容。"
    file = ("test.md", io.BytesIO(content.encode("utf-8")), "text/markdown")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 200, res.text
    data = res.json()["data"]
    assert data["import_status"] == "success"
    assert data["file_name"] == "test.md"
    assert data["file_type"] == ".md"
    assert data["page_id"] is not None
    assert data["knowledge_base_id"] == kb["id"]

    # 验证页面内容
    page_res = client.get(f"/api/v1/pages/{data['page_id']}")
    assert page_res.status_code == 200
    page = page_res.json()["data"]
    assert page["title"] == "测试标题"
    assert page["content"] == content


def test_upload_markdown_no_heading(client: TestClient) -> None:
    """无一级标题的 .md 文件应使用文件名作为标题。"""
    kb = _create_kb(client)
    content = "纯文本内容，无标题。"
    file = ("notes.md", io.BytesIO(content.encode("utf-8")), "text/markdown")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 200
    page = client.get(f"/api/v1/pages/{res.json()['data']['page_id']}")
    assert page.json()["data"]["title"] == "notes"


# ---------- TXT 导入 ----------


def test_upload_txt(client: TestClient) -> None:
    """上传 .txt 文件应使用文件名作为标题。"""
    kb = _create_kb(client)
    content = "普通文本内容\n第二行。"
    file = ("readme.txt", io.BytesIO(content.encode("utf-8")), "text/plain")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["import_status"] == "success"
    assert data["file_type"] == ".txt"

    page = client.get(f"/api/v1/pages/{data['page_id']}")
    assert page.json()["data"]["title"] == "readme"
    assert page.json()["data"]["content"] == content


# ---------- 文件校验 ----------


def test_upload_invalid_extension(client: TestClient) -> None:
    """不支持的文件类型应返回 400。"""
    kb = _create_kb(client)
    file = ("image.png", io.BytesIO(b"fake image"), "image/png")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 400


def test_upload_empty_file(client: TestClient) -> None:
    """空文件应能导入（无内容页面）。"""
    kb = _create_kb(client)
    file = ("empty.md", io.BytesIO(b""), "text/markdown")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["import_status"] == "success"


def test_upload_kb_not_found(client: TestClient) -> None:
    """不存在的知识库应返回 404。"""
    file = ("test.md", io.BytesIO(b"# title"), "text/markdown")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": 9999},
        files={"file": file},
    )
    assert res.status_code == 404


def test_upload_filename_sanitized(client: TestClient) -> None:
    """含路径穿越字符的文件名应被清理。"""
    kb = _create_kb(client)
    # 模拟穿越路径的文件名
    file = ("../../etc/test.md", io.BytesIO(b"# safe"), "text/markdown")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    # 文件名不应含 ../
    assert ".." not in data["file_name"]


# ---------- DOCX 导入 ----------


def test_upload_docx(client: TestClient) -> None:
    """上传 .docx 文件应成功导入并保留标题层级。"""
    from docx import Document as DocxDocument

    kb = _create_kb(client)

    # 用 python-docx 构建最小 docx
    doc = DocxDocument()
    doc.add_heading("Word 标题", level=1)
    doc.add_paragraph("第一段内容。")
    doc.add_heading("二级标题", level=2)
    doc.add_paragraph("第二段内容。")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    file = ("report.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 200, res.text
    data = res.json()["data"]
    assert data["import_status"] == "success"
    assert data["file_type"] == ".docx"
    assert data["page_id"] is not None

    # 验证页面标题来自 docx 第一个 Heading
    page_res = client.get(f"/api/v1/pages/{data['page_id']}")
    assert page_res.status_code == 200
    page = page_res.json()["data"]
    assert page["title"] == "Word 标题"


def test_upload_docx_no_heading(client: TestClient) -> None:
    """无标题样式的 .docx 文件应使用文件名作为标题。"""
    from docx import Document as DocxDocument

    kb = _create_kb(client)

    doc = DocxDocument()
    doc.add_paragraph("纯文字段落，无标题。")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    file = ("plain.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["import_status"] == "success"
    assert data["file_type"] == ".docx"

    page_res = client.get(f"/api/v1/pages/{data['page_id']}")
    assert page_res.status_code == 200
    page = page_res.json()["data"]
    assert page["title"] == "plain"


# ---------- 文件大小限制 ----------


def test_upload_file_too_large(client: TestClient) -> None:
    """超过 10MB 的文件应被拒绝。"""
    kb = _create_kb(client)
    # 生成超过 10MB 的数据
    large_content = b"x" * (10 * 1024 * 1024 + 1)
    file = ("large.md", io.BytesIO(large_content), "text/markdown")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 400


# ---------- 损坏文件 ----------


def test_upload_corrupt_docx(client: TestClient) -> None:
    """上传非有效 .docx 文件应返回解析失败错误。"""
    kb = _create_kb(client)
    # 伪造的 docx（不是有效的 zip/docx 格式）
    file = ("corrupt.docx", io.BytesIO(b"not a real docx file"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 500, res.text


def test_upload_utf8_encoding(client: TestClient) -> None:
    """UTF-8 编码的中文 .md 文件应正确导入。"""
    kb = _create_kb(client)
    content = "# 中文标题\n\n包含中文字符和特殊符号：①②③、简体繁體。"
    file = ("中文.md", io.BytesIO(content.encode("utf-8")), "text/markdown")

    res = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": file},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["import_status"] == "success"

    page_res = client.get(f"/api/v1/pages/{data['page_id']}")
    page = page_res.json()["data"]
    assert page["title"] == "中文标题"
    assert "①②③" in page["content"]